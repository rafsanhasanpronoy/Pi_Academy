from datetime import date, timedelta
from io import BytesIO

import openpyxl
from django.contrib import admin
from django.core.cache import cache
from django.test import TestCase, override_settings
from django.urls import reverse

from .admin import ExamResultAdmin, StudentAdmin, _finance_summary
from .forms import AdmissionInquiryForm, ContactMessageForm, ExpenseForm, StudentLookupForm
from .models import (
    AdmissionApplication, AdmissionInfo, Attendance, Batch, Class, Exam, ExamResult,
    Expense, Faculty, Student,
    StudentPayment, StudentPaymentReceipt, Subject, TeacherSalary,
)


def _make_xlsx(headers, rows):
    """Builds an in-memory .xlsx matching what a real upload looks like,
    so these tests exercise the same openpyxl parsing path as a real
    file — not a shortcut that skips it."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


class ContactMessageFormTests(TestCase):
    def test_rejects_when_no_phone_and_no_email(self):
        form = ContactMessageForm(data={
            "name": "Rafsan",
            "phone": "",
            "email": "",
            "message": "Hello",
        })
        self.assertFalse(form.is_valid())
        self.assertIn("__all__", form.errors)

    def test_valid_with_phone_only(self):
        form = ContactMessageForm(data={
            "name": "Rafsan",
            "phone": "01700000000",
            "email": "",
            "message": "Hello",
        })
        self.assertTrue(form.is_valid())

    def test_honeypot_rejects_bot_submission(self):
        form = ContactMessageForm(data={
            "name": "Bot",
            "phone": "01700000000",
            "email": "",
            "message": "spam",
            "website": "http://spam.example",  # a human never fills this in
        })
        self.assertFalse(form.is_valid())


class AdmissionInquiryFormTests(TestCase):
    def test_minimal_valid_submission(self):
        form = AdmissionInquiryForm(data={
            "student_name": "Test Student",
            "student_phone": "01700000000",
            "guardian_name": "",
            "guardian_phone": "",
            "class_obj": "",
            "message": "",
        })
        self.assertTrue(form.is_valid())

    def test_honeypot_rejects_bot_submission(self):
        form = AdmissionInquiryForm(data={
            "student_name": "Bot Student",
            "student_phone": "01700000000",
            "website": "http://spam.example",
        })
        self.assertFalse(form.is_valid())


class StudentResultsRateLimitTests(TestCase):
    """
    Covers the fix for the brute-force lookup risk flagged in QA: after
    the rate limit is hit, further POSTs must be blocked rather than
    silently allowed through.
    """

    def setUp(self):
        cache.clear()
        self.url = reverse("results")
        self.klass = Class.objects.create(class_name="Class 9", academic_year=2026)
        self.batch = Batch.objects.create(class_obj=self.klass, batch_name="Morning")
        self.student = Student.objects.create(
            student_code="STU-2026-001",
            class_obj=self.klass,
            batch=self.batch,
            full_name="Test Student",
            student_phone="01700000000",
        )

    @override_settings(RATELIMIT_ENABLE=True)
    def test_blocked_after_rate_limit_exceeded(self):
        payload = {"student_code": "WRONG-CODE", "student_phone": "0000000000"}
        responses = [self.client.post(self.url, payload) for _ in range(6)]
        # The 6th attempt within a minute from the same client should be
        # blocked (django-ratelimit raises Ratelimited -> 403 by default).
        self.assertEqual(responses[-1].status_code, 403)

    def test_correct_credentials_return_student_data(self):
        response = self.client.post(self.url, {
            "student_code": "STU-2026-001",
            "student_phone": "01700000000",
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Test Student")

    def test_wrong_credentials_do_not_leak_which_field_was_wrong(self):
        response = self.client.post(self.url, {
            "student_code": "STU-2026-001",
            "student_phone": "09999999999",
        })
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "Test Student")


class DashboardOrderingTests(TestCase):
    """Achievements with no date shouldn't unpredictably jump to the top."""

    def test_null_dates_sort_last(self):
        from .models import Achievement

        Achievement.objects.create(title="Dated", is_published=True, achievement_date="2026-01-01")
        Achievement.objects.create(title="Undated", is_published=True, achievement_date=None)

        response = self.client.get(reverse("achievements"))
        content = response.content.decode()
        self.assertLess(content.index("Dated"), content.index("Undated"))


class StudentBulkUploadTests(TestCase):
    """
    Covers StudentAdmin._process_bulk_upload — the spreadsheet-parsing
    path flagged in QA as the highest silent-failure risk in the admin,
    since a malformed row can otherwise fail quietly with nothing to
    catch it.
    """

    HEADERS = [
        "Full Name", "Class Name", "Academic Year", "Batch Name", "Gender",
        "Date of Birth (YYYY-MM-DD)", "Student Phone", "Father's Name",
        "Father's Phone", "Mother's Name", "Mother's Phone", "Address",
        "Admission Date (YYYY-MM-DD)", "Status",
    ]

    def setUp(self):
        self.admin = StudentAdmin(Student, admin.site)
        self.klass = Class.objects.create(class_name="Class 6", academic_year=2026)
        self.batch = Batch.objects.create(class_obj=self.klass, batch_name="Morning")

    def test_valid_row_creates_student_with_generated_code(self):
        upload = _make_xlsx(self.HEADERS, [[
            "Jane Doe", "Class 6", 2026, "Morning", "Female",
            "2014-05-12", "01700000000", "John Doe", "01700000001",
            "Mary Doe", "01700000002", "House 12, Road 4", "2026-01-10", "Active",
        ]])
        result = self.admin._process_bulk_upload(upload)
        self.assertEqual(result["created_count"], 1)
        self.assertEqual(result["errors"], [])
        student = Student.objects.get(full_name="Jane Doe")
        self.assertTrue(student.student_code.startswith("PiC6"))

    def test_missing_full_name_is_skipped_with_error_not_silently(self):
        upload = _make_xlsx(self.HEADERS, [[
            "", "Class 6", 2026, "Morning", "Female",
            "2014-05-12", "01700000000", "", "", "", "", "", "", "Active",
        ]])
        result = self.admin._process_bulk_upload(upload)
        self.assertEqual(result["created_count"], 0)
        self.assertEqual(Student.objects.count(), 0)
        self.assertTrue(any("missing Full Name" in e for e in result["errors"]))

    def test_unknown_class_is_skipped_with_error(self):
        upload = _make_xlsx(self.HEADERS, [[
            "Jane Doe", "Class 99", 2026, "", "Female",
            "", "", "", "", "", "", "", "", "Active",
        ]])
        result = self.admin._process_bulk_upload(upload)
        self.assertEqual(result["created_count"], 0)
        self.assertTrue(any("no class matching" in e for e in result["errors"]))

    def test_unknown_batch_is_skipped_with_error(self):
        upload = _make_xlsx(self.HEADERS, [[
            "Jane Doe", "Class 6", 2026, "Evening (doesn't exist)", "Female",
            "", "", "", "", "", "", "", "", "Active",
        ]])
        result = self.admin._process_bulk_upload(upload)
        self.assertEqual(result["created_count"], 0)
        self.assertTrue(any("no batch matching" in e for e in result["errors"]))

    def test_blank_rows_are_skipped_without_error(self):
        upload = _make_xlsx(self.HEADERS, [
            [None] * len(self.HEADERS),
            ["Jane Doe", "Class 6", 2026, "Morning", "Female",
             "", "", "", "", "", "", "", "", "Active"],
        ])
        result = self.admin._process_bulk_upload(upload)
        self.assertEqual(result["created_count"], 1)
        self.assertEqual(result["errors"], [])

    def test_wrong_headers_are_rejected_before_touching_any_row(self):
        upload = _make_xlsx(["Wrong", "Headers"], [["a", "b"]])
        result = self.admin._process_bulk_upload(upload)
        self.assertEqual(result["created_count"], 0)
        self.assertTrue(any("don't match the template" in e for e in result["errors"]))

    def test_not_an_excel_file_fails_gracefully(self):
        garbage = BytesIO(b"this is not a real xlsx file")
        result = self.admin._process_bulk_upload(garbage)
        self.assertEqual(result["created_count"], 0)
        self.assertTrue(result["errors"])  # a friendly message, not a crash

    def test_sequential_rows_in_same_class_and_batch_get_distinct_codes(self):
        upload = _make_xlsx(self.HEADERS, [
            ["Student A", "Class 6", 2026, "Morning", "", "", "", "", "", "", "", "", "", "Active"],
            ["Student B", "Class 6", 2026, "Morning", "", "", "", "", "", "", "", "", "", "Active"],
        ])
        result = self.admin._process_bulk_upload(upload)
        self.assertEqual(result["created_count"], 2)
        codes = set(Student.objects.values_list("student_code", flat=True))
        self.assertEqual(len(codes), 2)  # no collision between the two rows


class ExamResultBulkUploadTests(TestCase):
    """
    Covers ExamResultAdmin._process_bulk_upload — the mark-sheet parser,
    including the update_or_create re-upload behavior that's easy to
    accidentally break into duplicate rows or silent overwrites.
    """

    def setUp(self):
        self.admin = ExamResultAdmin(ExamResult, admin.site)
        self.klass = Class.objects.create(class_name="Class 7", academic_year=2026)
        self.batch = Batch.objects.create(class_obj=self.klass, batch_name="Morning")
        self.subject = Subject.objects.create(class_obj=self.klass, subject_name="Physics")
        self.exam = Exam.objects.create(
            class_obj=self.klass, exam_name="Midterm", exam_date="2026-06-01"
        )
        self.student = Student.objects.create(
            student_code="PiC7B1001", class_obj=self.klass, batch=self.batch,
            full_name="Test Student",
        )

    def test_valid_marks_create_exam_result(self):
        upload = _make_xlsx(
            ["Student ID", "Full Name", "Physics"],
            [["PiC7B1001", "Test Student", 85]],
        )
        result = self.admin._process_bulk_upload(upload, self.exam)
        self.assertEqual(result["saved_count"], 1)
        self.assertEqual(result["errors"], [])
        er = ExamResult.objects.get(exam=self.exam, student=self.student, subject=self.subject)
        self.assertEqual(er.marks, 85)

    def test_reuploading_corrected_marks_updates_not_duplicates(self):
        headers = ["Student ID", "Full Name", "Physics"]
        first = _make_xlsx(headers, [["PiC7B1001", "Test Student", 60]])
        self.admin._process_bulk_upload(first, self.exam)

        corrected = _make_xlsx(headers, [["PiC7B1001", "Test Student", 92]])
        result = self.admin._process_bulk_upload(corrected, self.exam)

        self.assertEqual(result["saved_count"], 1)
        self.assertEqual(
            ExamResult.objects.filter(exam=self.exam, student=self.student, subject=self.subject).count(),
            1,  # updated in place, not a second row
        )
        er = ExamResult.objects.get(exam=self.exam, student=self.student, subject=self.subject)
        self.assertEqual(er.marks, 92)

    def test_unrecognized_subject_column_is_flagged_not_silently_dropped(self):
        upload = _make_xlsx(
            ["Student ID", "Full Name", "Chemistry"],  # not in this class
            [["PiC7B1001", "Test Student", 70]],
        )
        result = self.admin._process_bulk_upload(upload, self.exam)
        self.assertEqual(result["saved_count"], 0)
        self.assertTrue(any("Chemistry" in e for e in result["errors"]))

    def test_non_numeric_marks_are_skipped_with_error(self):
        upload = _make_xlsx(
            ["Student ID", "Full Name", "Physics"],
            [["PiC7B1001", "Test Student", "absent"]],
        )
        result = self.admin._process_bulk_upload(upload, self.exam)
        self.assertEqual(result["saved_count"], 0)
        self.assertTrue(any("isn't a number" in e for e in result["errors"]))

    def test_negative_marks_are_rejected(self):
        upload = _make_xlsx(
            ["Student ID", "Full Name", "Physics"],
            [["PiC7B1001", "Test Student", -5]],
        )
        result = self.admin._process_bulk_upload(upload, self.exam)
        self.assertEqual(result["saved_count"], 0)
        self.assertTrue(any("negative" in e for e in result["errors"]))

    def test_unknown_student_id_is_skipped_with_error(self):
        upload = _make_xlsx(
            ["Student ID", "Full Name", "Physics"],
            [["DOES-NOT-EXIST", "Nobody", 70]],
        )
        result = self.admin._process_bulk_upload(upload, self.exam)
        self.assertEqual(result["saved_count"], 0)
        self.assertTrue(any("no student" in e for e in result["errors"]))


class ExpenseFormTests(TestCase):
    def _data(self, **overrides):
        data = {
            "title": "Office rent",
            "category": "Rent",
            "amount": "12000",
            "expense_date": date.today().isoformat(),
            "payment_method": "Cash",
            "remarks": "",
        }
        data.update(overrides)
        return data

    def test_valid_expense(self):
        self.assertTrue(ExpenseForm(data=self._data()).is_valid())

    def test_zero_and_negative_amounts_rejected(self):
        self.assertFalse(ExpenseForm(data=self._data(amount="0")).is_valid())
        self.assertFalse(ExpenseForm(data=self._data(amount="-50")).is_valid())

    def test_future_date_rejected(self):
        future = (date.today() + timedelta(days=3)).isoformat()
        self.assertFalse(ExpenseForm(data=self._data(expense_date=future)).is_valid())

    def test_amount_is_rounded_to_whole_taka(self):
        form = ExpenseForm(data=self._data(amount="150.40"))
        self.assertTrue(form.is_valid())
        self.assertEqual(form.cleaned_data["amount"], 150)


class FinanceSummaryTests(TestCase):
    """Income = receipts (minus discount) + legacy payments.
    Expenses = salaries actually PAID + other expenses. Net = difference."""

    def setUp(self):
        klass = Class.objects.create(class_name="Class 6", academic_year=2026)
        self.student = Student.objects.create(
            student_code="PiC6B0001", class_obj=klass, full_name="Test Student"
        )
        self.teacher = Faculty.objects.create(full_name="Test Teacher")
        self.today = date.today()

    def _receipt(self, total, discount, when=None, number=None):
        return StudentPaymentReceipt.objects.create(
            student=self.student, payment_date=when or self.today,
            items=[], total_amount=total, total_discount=discount,
            receipt_number=number,
        )

    def test_income_subtracts_receipt_discount(self):
        self._receipt(3000, 500, number="PAY-T-1")
        self.assertEqual(_finance_summary()["income"], 2500)

    def test_legacy_payments_count_as_income(self):
        StudentPayment.objects.create(
            student=self.student, payment_type="Monthly Fee",
            amount=1000, payment_date=self.today,
        )
        self._receipt(2000, 0, number="PAY-T-2")
        self.assertEqual(_finance_summary()["income"], 3000)

    def test_only_paid_salary_counts_as_expense(self):
        TeacherSalary.objects.create(
            teacher=self.teacher, salary_month=self.today.replace(day=1),
            amount=6000, net_salary=6000, paid_amount=4000, due_amount=2000,
            payment_date=self.today,
        )
        summary = _finance_summary()
        self.assertEqual(summary["salaries_paid"], 4000)

    def test_unpaid_salary_is_not_an_expense(self):
        TeacherSalary.objects.create(
            teacher=self.teacher, salary_month=self.today.replace(day=1),
            amount=6000, net_salary=6000, paid_amount=0, due_amount=6000,
        )
        self.assertEqual(_finance_summary()["salaries_paid"], 0)

    def test_net_is_income_minus_all_expenses(self):
        self._receipt(10000, 1000, number="PAY-T-3")  # income 9000
        TeacherSalary.objects.create(
            teacher=self.teacher, salary_month=self.today.replace(day=1),
            amount=3000, net_salary=3000, paid_amount=3000, due_amount=0,
            payment_date=self.today,
        )
        Expense.objects.create(
            title="Rent", category="Rent", amount=2000, expense_date=self.today,
        )
        summary = _finance_summary()
        self.assertEqual(summary["income"], 9000)
        self.assertEqual(summary["expenses"], 5000)
        self.assertEqual(summary["net"], 4000)

    def test_net_can_go_negative(self):
        Expense.objects.create(
            title="Repairs", category="Maintenance", amount=500, expense_date=self.today,
        )
        self.assertEqual(_finance_summary()["net"], -500)

    def test_date_range_excludes_other_months(self):
        old = self.today - timedelta(days=90)
        self._receipt(5000, 0, when=old, number="PAY-T-4")
        Expense.objects.create(title="Old", category="Other", amount=700, expense_date=old)
        window = _finance_summary(self.today - timedelta(days=30), self.today)
        self.assertEqual(window["income"], 0)
        self.assertEqual(window["expenses"], 0)


class ExpenseAdminTests(TestCase):
    def setUp(self):
        from django.contrib.auth import get_user_model
        self.admin_user = get_user_model().objects.create_superuser(
            "boss", "boss@example.com", "pw12345!"
        )
        self.client.force_login(self.admin_user)

    def test_add_expense_through_admin(self):
        response = self.client.post(reverse("admin:core_expense_add"), {
            "title": "Markers", "category": "Supplies", "amount": "250",
            "expense_date": date.today().isoformat(), "payment_method": "Cash",
            "remarks": "",
        })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Expense.objects.get().amount, 250)

    def test_list_and_delete(self):
        expense = Expense.objects.create(
            title="Tea", category="Events", amount=90, expense_date=date.today(),
        )
        self.assertEqual(self.client.get(reverse("admin:core_expense_changelist")).status_code, 200)
        self.client.post(reverse("admin:core_expense_delete", args=[expense.pk]))
        self.assertFalse(Expense.objects.exists())

    def test_dashboard_shows_income_expenses_and_net(self):
        Expense.objects.create(
            title="Rent", category="Rent", amount=800, expense_date=date.today(),
        )
        response = self.client.get(reverse("admin:index"))
        self.assertEqual(response.status_code, 200)
        finance = response.context["dashboard_finance"]
        self.assertEqual(finance["month"]["expenses"], 800)
        self.assertEqual(finance["month"]["net"], -800)
        self.assertContains(response, "Net revenue")


class StudentAdmissionFieldsTests(TestCase):
    def setUp(self):
        self.klass = Class.objects.create(class_name="Class 8", academic_year=2026)

    def _data(self, **overrides):
        data = {
            "class_obj": self.klass.pk, "full_name": "Rahim Uddin",
            "date_of_birth": "2012-03-10", "admission_date": date.today().isoformat(),
            "institution_name": "Dhaka Residential School", "blood_group": "O+",
            "reference": "", "status": "Active",
        }
        data.update(overrides)
        return data

    def test_valid_new_admission_with_reference_left_blank(self):
        from .forms import StudentForm
        self.assertTrue(StudentForm(data=self._data()).is_valid())

    def test_new_admission_requires_dob_admission_date_school_and_blood_group(self):
        from .forms import StudentForm
        for field in ("date_of_birth", "admission_date", "institution_name", "blood_group"):
            form = StudentForm(data=self._data(**{field: ""}))
            self.assertFalse(form.is_valid(), field)
            self.assertIn(field, form.errors)

    def test_invalid_blood_group_rejected(self):
        from .forms import StudentForm
        self.assertFalse(StudentForm(data=self._data(blood_group="Z+")).is_valid())

    def test_legacy_student_without_new_fields_can_still_be_edited(self):
        from .forms import StudentForm
        legacy = Student.objects.create(
            student_code="PiC8B0001", class_obj=self.klass, full_name="Old Student",
        )
        form = StudentForm(
            data={"class_obj": self.klass.pk, "full_name": "Old Student", "status": "Left"},
            instance=legacy,
        )
        self.assertTrue(form.is_valid(), form.errors)

    def test_existing_value_cannot_be_blanked_out(self):
        from .forms import StudentForm
        student = Student.objects.create(
            student_code="PiC8B0002", class_obj=self.klass, full_name="Has Data",
            blood_group="A+", institution_name="X School",
            date_of_birth=date(2012, 1, 1), admission_date=date.today(),
        )
        form = StudentForm(
            data=self._data(full_name="Has Data", blood_group=""), instance=student,
        )
        self.assertFalse(form.is_valid())
        self.assertIn("blood_group", form.errors)


class StudentBulkUploadNewFieldsTests(TestCase):
    """Bulk upload with the School/College, Blood Group and Reference
    columns, plus the cell-cleaning fixes (phones, dates, status)."""

    from .admin import BULK_STUDENT_HEADERS as NEW_HEADERS  # 17 columns

    def setUp(self):
        self.admin = StudentAdmin(Student, admin.site)
        self.klass = Class.objects.create(class_name="Class 6", academic_year=2026)
        Batch.objects.create(class_obj=self.klass, batch_name="Morning")

    def _row(self, **o):
        row = {
            "name": "Jane Doe", "gender": "Female", "dob": "2014-05-12",
            "phone": "01700000000", "adm": "2026-01-10", "status": "Active",
            "school": "Dhaka Residential School", "blood": "O+", "ref": "",
        }
        row.update(o)
        return [
            row["name"], "Class 6", 2026, "Morning", row["gender"], row["dob"],
            row["phone"], "", "", "", "", "", row["adm"], row["status"],
            row["school"], row["blood"], row["ref"],
        ]

    def _run(self, *rows):
        return self.admin._process_bulk_upload(_make_xlsx(self.NEW_HEADERS, list(rows)))

    def test_template_has_seventeen_columns_new_ones_last(self):
        self.assertEqual(len(self.NEW_HEADERS), 17)
        self.assertEqual(self.NEW_HEADERS[14:], [
            "School / College Name", "Blood Group", "Reference (optional)",
        ])

    def test_new_fields_are_saved(self):
        result = self._run(self._row(ref="Karim sir"))
        self.assertEqual(result["errors"], [])
        s = Student.objects.get(full_name="Jane Doe")
        self.assertEqual((s.institution_name, s.blood_group, s.reference),
                         ("Dhaka Residential School", "O+", "Karim sir"))

    def test_reference_is_optional(self):
        self.assertEqual(self._run(self._row(ref=""))["created_count"], 1)

    def test_blood_group_is_normalized(self):
        self._run(self._row(blood=" ab - "))
        self.assertEqual(Student.objects.get().blood_group, "AB-")

    def test_unknown_blood_group_is_skipped_with_error(self):
        result = self._run(self._row(blood="Z+"))
        self.assertEqual(result["created_count"], 0)
        self.assertTrue(any("Blood Group" in e for e in result["errors"]))

    def test_required_fields_enforced_with_new_template(self):
        for field, label in (("dob", "Date of Birth"), ("adm", "Admission Date"),
                             ("school", "School / College Name"), ("blood", "Blood Group")):
            Student.objects.all().delete()
            result = self._run(self._row(**{field: ""}))
            self.assertEqual(result["created_count"], 0, field)
            self.assertTrue(any(label in e for e in result["errors"]), field)

    def test_old_14_column_file_still_accepted_without_new_fields(self):
        legacy = StudentBulkUploadTests.HEADERS
        upload = _make_xlsx(legacy, [[
            "Old Style", "Class 6", 2026, "Morning", "", "", "", "", "", "", "", "", "", "Active",
        ]])
        result = self.admin._process_bulk_upload(upload)
        self.assertEqual(result["created_count"], 1)
        self.assertEqual(result["errors"], [])

    def test_partial_or_misnamed_new_columns_rejected(self):
        headers = StudentBulkUploadTests.HEADERS + ["School / College Name", "Wrong", "Reference (optional)"]
        result = self.admin._process_bulk_upload(_make_xlsx(headers, []))
        self.assertEqual(result["created_count"], 0)
        self.assertTrue(any("new columns" in e for e in result["errors"]))

    def test_numeric_phone_cells_get_their_leading_zero_back(self):
        self._run(self._row(phone=1700000000), )
        self.assertEqual(Student.objects.get().student_phone, "01700000000")
        Student.objects.all().delete()
        self._run(self._row(phone=1700000000.0))
        self.assertEqual(Student.objects.get().student_phone, "01700000000")

    def test_unreadable_date_is_flagged_not_dropped(self):
        result = self._run(self._row(dob="12th May"))
        self.assertEqual(result["created_count"], 0)
        self.assertTrue(any("Date of Birth" in e for e in result["errors"]))

    def test_unknown_status_and_gender_are_flagged_not_defaulted(self):
        result = self._run(self._row(status="Enrolled"), self._row(name="B", gender="Robot"))
        self.assertEqual(result["created_count"], 0)
        self.assertTrue(any("Status" in e for e in result["errors"]))
        self.assertTrue(any("Gender" in e for e in result["errors"]))

    def test_downloaded_template_matches_upload_format(self):
        from django.contrib.auth import get_user_model
        user = get_user_model().objects.create_superuser("boss", "b@x.com", "pw12345!")
        self.client.force_login(user)
        response = self.client.get(reverse("admin:core_student_bulk_template"))
        self.assertEqual(response.status_code, 200)
        sheet = openpyxl.load_workbook(BytesIO(response.content)).active
        self.assertEqual([c.value for c in sheet[1]], self.NEW_HEADERS)
        # the sample row in the template must itself upload cleanly
        result = self.admin._process_bulk_upload(BytesIO(response.content))
        self.assertEqual(result["errors"], [])
        self.assertEqual(result["created_count"], 1)


class ApplyNowPublicTests(TestCase):
    """The public Apply Now flow: collects an application, never admits."""

    def setUp(self):
        self.course = AdmissionInfo.objects.create(title="SSC Science Batch", is_active=True)
        self.closed = AdmissionInfo.objects.create(title="Old Course", is_active=False)

    def _data(self, **o):
        data = {
            "course": self.course.pk, "preferred_batch": "Morning",
            "student_name": "Rahim Uddin", "date_of_birth": "2010-04-02", "gender": "Male",
            "institution_name": "Narayanganj High School", "blood_group": "B+",
            "student_phone": "", "address": "Narayanganj",
            "guardian_name": "Karim Uddin", "guardian_relation": "Father",
            "guardian_phone": "01711111111", "reference": "", "message": "",
            "understand": "on",
        }
        data.update(o)
        return data

    def test_apply_page_renders_and_explains_verification(self):
        response = self.client.get(reverse("apply"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "How admission works")
        self.assertContains(response, "SSC Science Batch")
        self.assertNotContains(response, "Old Course")  # inactive courses aren't offered

    def test_course_query_param_preselects_the_course(self):
        response = self.client.get(reverse("apply") + f"?course={self.course.pk}")
        self.assertEqual(response.context["form"].initial["course"], self.course.pk)

    def test_inactive_or_junk_course_param_is_ignored(self):
        for value in (str(self.closed.pk), "abc", "99999"):
            response = self.client.get(reverse("apply") + f"?course={value}")
            self.assertNotIn("course", response.context["form"].initial)

    def test_submit_creates_pending_application_and_no_student(self):
        response = self.client.post(reverse("apply"), self._data())
        self.assertRedirects(response, reverse("apply_done"))
        app = AdmissionApplication.objects.get()
        self.assertEqual(app.status, "Pending")
        self.assertEqual(app.course_title, "SSC Science Batch")
        self.assertEqual(app.guardian_phone, "01711111111")
        self.assertFalse(Student.objects.exists())  # applying never admits

    def test_public_cannot_set_status(self):
        self.client.post(reverse("apply"), self._data(status="Admitted"))
        self.assertEqual(AdmissionApplication.objects.get().status, "Pending")

    def test_done_page_shows_reference_and_not_admitted_notice(self):
        self.client.post(reverse("apply"), self._data())
        response = self.client.get(reverse("apply_done"))
        ref = AdmissionApplication.objects.get().reference_code
        self.assertContains(response, ref)
        self.assertContains(response, "not a confirmed admission")
        self.assertContains(response, "guardian")

    def test_done_page_without_an_application_redirects_back(self):
        self.assertRedirects(self.client.get(reverse("apply_done")), reverse("apply"))

    def test_must_acknowledge_verification(self):
        data = self._data(); data.pop("understand")
        response = self.client.post(reverse("apply"), data)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(AdmissionApplication.objects.exists())

    def test_required_fields(self):
        for field in ("course", "student_name", "date_of_birth", "gender",
                      "institution_name", "guardian_name", "guardian_phone"):
            response = self.client.post(reverse("apply"), self._data(**{field: ""}))
            self.assertEqual(response.status_code, 200, field)
            self.assertIn(field, response.context["form"].errors, field)
        self.assertFalse(AdmissionApplication.objects.exists())

    def test_optional_fields_can_be_blank(self):
        data = self._data(blood_group="", student_phone="", address="", reference="",
                          message="", preferred_batch="")
        self.assertEqual(self.client.post(reverse("apply"), data).status_code, 302)

    def test_bad_values_rejected(self):
        future = (date.today() + timedelta(days=5)).isoformat()
        for overrides, field in (
            ({"date_of_birth": future}, "date_of_birth"),
            ({"guardian_phone": "abc"}, "guardian_phone"),
            ({"guardian_phone": "123"}, "guardian_phone"),
            ({"student_phone": "xx"}, "student_phone"),
            ({"blood_group": "Z+"}, "blood_group"),
            ({"course": self.closed.pk}, "course"),
        ):
            response = self.client.post(reverse("apply"), self._data(**overrides))
            self.assertEqual(response.status_code, 200, overrides)
            self.assertIn(field, response.context["form"].errors, overrides)

    def test_honeypot_blocks_bots(self):
        response = self.client.post(reverse("apply"), self._data(website="http://spam.example"))
        self.assertEqual(response.status_code, 200)
        self.assertFalse(AdmissionApplication.objects.exists())

    def test_duplicate_open_application_is_blocked(self):
        self.client.post(reverse("apply"), self._data())
        response = self.client.post(reverse("apply"), self._data())
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "already have an open application")
        self.assertEqual(AdmissionApplication.objects.count(), 1)

    def test_rejected_application_does_not_block_reapplying(self):
        self.client.post(reverse("apply"), self._data())
        AdmissionApplication.objects.update(status="Rejected")
        self.assertEqual(self.client.post(reverse("apply"), self._data()).status_code, 302)

    def test_no_active_courses_shows_closed_message(self):
        AdmissionInfo.objects.update(is_active=False)
        response = self.client.get(reverse("apply"))
        self.assertContains(response, "aren't open right now")

    def test_apply_now_buttons_point_to_the_application_form(self):
        apply_url = reverse("apply")
        for name in ("home", "classes", "admission"):
            response = self.client.get(reverse(name))
            self.assertContains(response, f'href="{apply_url}"', msg_prefix=name)

    def test_course_popup_apply_button_carries_the_course(self):
        response = self.client.get(reverse("admission"))
        self.assertContains(response, f'data-id="{self.course.pk}"')
        self.assertContains(response, 'data-apply-url="/apply/"')
        # the inquiry form is still there, separate from Apply Now
        self.assertContains(response, "Send an inquiry")


class ApplicationStaffFlowTests(TestCase):
    """Staff review an application, verify it, then admit — the Student is
    only created by staff, and the application is marked Admitted then."""

    def setUp(self):
        from django.contrib.auth import get_user_model
        self.user = get_user_model().objects.create_superuser("boss", "b@x.com", "pw12345!")
        self.client.force_login(self.user)
        self.klass = Class.objects.create(class_name="Class 9", academic_year=2026)
        self.course = AdmissionInfo.objects.create(title="SSC Science Batch", is_active=True)
        self.app = AdmissionApplication.objects.create(
            course=self.course, course_title="SSC Science Batch",
            student_name="Rahim Uddin", date_of_birth=date(2010, 4, 2), gender="Male",
            institution_name="Narayanganj High School", blood_group="B+",
            address="Narayanganj", guardian_name="Karim Uddin",
            guardian_relation="Father", guardian_phone="01711111111",
            reference="Sir Hasan", status="Pending",
        )

    def test_reference_code_format(self):
        self.assertRegex(self.app.reference_code, r"^APP-\d{4}-\d{5}$")

    def test_list_detail_and_edit_pages_render(self):
        for name, args in (("changelist", []), ("view", [self.app.pk]),
                           ("change", [self.app.pk]), ("add", [])):
            response = self.client.get(reverse(f"admin:core_admissionapplication_{name}", args=args))
            self.assertEqual(response.status_code, 200, name)
        detail = self.client.get(reverse("admin:core_admissionapplication_view", args=[self.app.pk]))
        self.assertContains(detail, "Verify before admitting")
        self.assertContains(detail, "01711111111")

    def test_list_search_and_status_filter(self):
        url = reverse("admin:core_admissionapplication_changelist")
        self.assertContains(self.client.get(url + "?q=Rahim"), "Rahim Uddin")
        self.assertNotContains(self.client.get(url + "?q=Nobody"), "Rahim Uddin")
        self.assertNotContains(self.client.get(url + "?status=Verified"), "Rahim Uddin")

    def test_quick_status_change(self):
        url = reverse("admin:core_admissionapplication_set_status", args=[self.app.pk])
        self.client.post(url, {"status": "Contacted"})
        self.app.refresh_from_db()
        self.assertEqual(self.app.status, "Contacted")

    def test_quick_status_cannot_mark_admitted(self):
        url = reverse("admin:core_admissionapplication_set_status", args=[self.app.pk])
        self.client.post(url, {"status": "Admitted"})
        self.app.refresh_from_db()
        self.assertEqual(self.app.status, "Pending")

    def test_status_change_requires_post(self):
        url = reverse("admin:core_admissionapplication_set_status", args=[self.app.pk])
        self.client.get(url + "?status=Verified")
        self.app.refresh_from_db()
        self.assertEqual(self.app.status, "Pending")

    def test_start_admission_opens_prefilled_student_form(self):
        response = self.client.get(
            reverse("admin:core_admissionapplication_start_admission", args=[self.app.pk]),
            follow=True,
        )
        self.assertEqual(response.status_code, 200)
        form = response.context["form"]
        self.assertEqual(form.initial["full_name"], "Rahim Uddin")
        self.assertEqual(form.initial["father_name"], "Karim Uddin")
        self.assertEqual(form.initial["father_phone"], "01711111111")
        self.assertEqual(form.initial["institution_name"], "Narayanganj High School")
        self.assertEqual(form.initial["reference"], "Sir Hasan")
        self.assertContains(response, "Admitting from application")

    def test_saving_the_student_marks_the_application_admitted(self):
        response = self.client.post(reverse("admin:core_student_add"), {
            "application": self.app.pk, "class_obj": self.klass.pk,
            "full_name": "Rahim Uddin", "date_of_birth": "2010-04-02",
            "admission_date": date.today().isoformat(),
            "institution_name": "Narayanganj High School", "blood_group": "B+",
            "status": "Active",
        })
        self.assertEqual(response.status_code, 302)
        student = Student.objects.get(full_name="Rahim Uddin")
        self.app.refresh_from_db()
        self.assertEqual(self.app.status, "Admitted")
        self.assertEqual(self.app.student_id, student.pk)

    def test_plain_student_add_does_not_touch_applications(self):
        self.client.post(reverse("admin:core_student_add"), {
            "class_obj": self.klass.pk, "full_name": "Someone Else",
            "date_of_birth": "2010-01-01", "admission_date": date.today().isoformat(),
            "institution_name": "X School", "blood_group": "O+", "status": "Active",
        })
        self.app.refresh_from_db()
        self.assertEqual(self.app.status, "Pending")

    def test_failed_student_save_leaves_application_pending(self):
        self.client.post(reverse("admin:core_student_add"), {
            "application": self.app.pk, "class_obj": self.klass.pk,
            "full_name": "Rahim Uddin", "status": "Active",  # missing required fields
        })
        self.app.refresh_from_db()
        self.assertEqual(self.app.status, "Pending")
        self.assertFalse(Student.objects.exists())

    def test_already_admitted_application_redirects_to_its_student(self):
        student = Student.objects.create(
            student_code="PiC9B0001", class_obj=self.klass, full_name="Rahim Uddin",
        )
        self.app.status, self.app.student = "Admitted", student
        self.app.save()
        response = self.client.get(
            reverse("admin:core_admissionapplication_start_admission", args=[self.app.pk])
        )
        self.assertRedirects(
            response, reverse("admin:core_student_view", args=[student.pk]),
            fetch_redirect_response=False,
        )

    def test_staff_can_edit_and_delete(self):
        self.client.post(
            reverse("admin:core_admissionapplication_change", args=[self.app.pk]),
            {"status": "Verified", "staff_notes": "Father visited", "course": self.course.pk,
             "student_name": "Rahim Uddin", "date_of_birth": "2010-04-02",
             "guardian_name": "Karim Uddin", "guardian_relation": "Father",
             "guardian_phone": "01711111111"},
        )
        self.app.refresh_from_db()
        self.assertEqual((self.app.status, self.app.staff_notes), ("Verified", "Father visited"))
        self.client.post(reverse("admin:core_admissionapplication_delete", args=[self.app.pk]))
        self.assertFalse(AdmissionApplication.objects.exists())

    def test_dashboard_counts_applications_to_verify(self):
        response = self.client.get(reverse("admin:index"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["dashboard_stats"]["pending_applications"], 1)
        self.assertContains(response, "Applications to verify")

    def test_anonymous_users_cannot_reach_application_admin(self):
        self.client.logout()
        for name, args in (("changelist", []), ("view", [self.app.pk])):
            response = self.client.get(reverse(f"admin:core_admissionapplication_{name}", args=args))
            self.assertEqual(response.status_code, 302)
            self.assertIn("login", response["Location"])


class GradingTests(TestCase):
    """core.results — the grade scale, GPA rules, and report building."""

    def test_grade_boundaries_out_of_100(self):
        from decimal import Decimal as D
        from .results import grade_for
        cases = {"100": "A+", "80": "A+", "79.99": "A", "70": "A", "69.99": "A-", "60": "A-",
                 "59.99": "B", "50": "B", "49.99": "C", "40": "C", "39.99": "D", "33": "D",
                 "32.99": "F", "0": "F"}
        for marks, expected in cases.items():
            self.assertEqual(grade_for(D(marks), D("100"))[0], expected, marks)

    def test_grade_uses_percentage_of_full_marks(self):
        from decimal import Decimal as D
        from .results import grade_for
        self.assertEqual(grade_for(D("40"), D("50")), ("A+", D("5.00")))   # 80%
        self.assertEqual(grade_for(D("16"), D("50"))[0], "F")             # 32%
        self.assertEqual(grade_for(D("16.5"), D("50"))[0], "D")           # 33%

    def test_zero_or_missing_full_marks_does_not_crash(self):
        from decimal import Decimal as D
        from .results import grade_for
        self.assertEqual(grade_for(D("50"), D("0"))[0], "F")
        self.assertEqual(grade_for(D("85"), None)[0], "A+")  # falls back to 100

    def _results(self, marks_by_subject, full_marks=100):
        klass = Class.objects.create(class_name="Class 8", academic_year=2026)
        exam = Exam.objects.create(class_obj=klass, exam_name="Term 1",
                                   exam_date="2026-05-01", full_marks=full_marks)
        student = Student.objects.create(student_code="PiC8B0001", class_obj=klass, full_name="Sam")
        for name, marks in marks_by_subject.items():
            subject = Subject.objects.create(class_obj=klass, subject_name=name)
            ExamResult.objects.create(exam=exam, student=student, subject=subject, marks=marks)
        return student

    def test_report_totals_gpa_and_pass(self):
        from .results import reports_for_student
        student = self._results({"Math": 90, "Physics": 75, "Chemistry": 65})  # 5 + 4 + 3.5
        report = reports_for_student(student)[0]
        self.assertEqual(report["total"], 230)
        self.assertEqual(report["full_total"], 300)
        self.assertEqual(str(report["percentage"]), "76.67")
        self.assertEqual(str(report["gpa"]), "4.17")
        self.assertEqual(report["overall_grade"], "A")
        self.assertTrue(report["passed"])

    def test_failing_any_subject_fails_the_exam_and_zeroes_gpa(self):
        from .results import reports_for_student
        student = self._results({"Math": 95, "Physics": 20})
        report = reports_for_student(student)[0]
        self.assertFalse(report["passed"])
        self.assertEqual(str(report["gpa"]), "0.00")
        self.assertEqual(report["overall_grade"], "F")
        self.assertEqual([r["passed"] for r in report["rows"]], [True, False])

    def test_subjects_are_listed_alphabetically_and_exams_newest_first(self):
        from .results import reports_for_student
        student = self._results({"Zoology": 50, "Algebra": 60})
        klass = student.class_obj
        newer = Exam.objects.create(class_obj=klass, exam_name="Term 2", exam_date="2026-09-01")
        ExamResult.objects.create(exam=newer, student=student,
                                  subject=Subject.objects.get(subject_name="Algebra"), marks=70)
        reports = reports_for_student(student)
        self.assertEqual([r["exam"].exam_name for r in reports], ["Term 2", "Term 1"])
        self.assertEqual([r["subject"] for r in reports[1]["rows"]], ["Algebra", "Zoology"])

    def test_hand_typed_grade_is_respected(self):
        from .results import reports_for_student
        student = self._results({"Math": 82})
        ExamResult.objects.update(grade="A")  # moderated down by staff
        row = reports_for_student(student)[0]["rows"][0]
        self.assertEqual(row["grade"], "A")
        self.assertEqual(str(row["grade_point"]), "4.00")

    def test_model_save_fills_missing_grade(self):
        student = self._results({"Math": 72})
        self.assertEqual(ExamResult.objects.get().grade, "A")

    def test_exam_full_marks_defaults_to_100(self):
        klass = Class.objects.create(class_name="C", academic_year=2026)
        self.assertEqual(Exam.objects.create(class_obj=klass, exam_name="E", exam_date="2026-01-01").full_marks, 100)


class ExamFullMarksFormTests(TestCase):
    def setUp(self):
        self.klass = Class.objects.create(class_name="Class 6", academic_year=2026)

    def _data(self, **o):
        d = {"class_obj": self.klass.pk, "exam_name": "Quiz", "exam_date": "2026-03-01", "full_marks": "50"}
        d.update(o)
        return d

    def test_valid_full_marks(self):
        from .forms import ExamForm
        self.assertTrue(ExamForm(data=self._data()).is_valid())

    def test_invalid_full_marks_rejected(self):
        from .forms import ExamForm
        for bad in ("0", "-5", "1000", "abc", ""):
            self.assertFalse(ExamForm(data=self._data(full_marks=bad)).is_valid(), bad)


class ExamResultFormGradingTests(TestCase):
    def setUp(self):
        self.klass = Class.objects.create(class_name="Class 6", academic_year=2026)
        self.exam = Exam.objects.create(class_obj=self.klass, exam_name="Quiz",
                                        exam_date="2026-03-01", full_marks=50)
        self.subject = Subject.objects.create(class_obj=self.klass, subject_name="Math")
        self.student = Student.objects.create(student_code="PiC6B0001", class_obj=self.klass, full_name="Ann")

    def _form(self, marks, grade="", instance=None):
        from .forms import ExamResultForm
        return ExamResultForm(
            data={"exam": self.exam.pk, "student": self.student.pk, "subject": self.subject.pk,
                  "marks": marks, "grade": grade},
            instance=instance, lock_exam=self.exam,
        )

    def test_blank_grade_is_calculated(self):
        form = self._form("42")  # 84% of 50
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["grade"], "A+")

    def test_hand_typed_grade_is_kept(self):
        form = self._form("42", grade="A")
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["grade"], "A")

    def test_marks_above_full_marks_rejected(self):
        form = self._form("51")
        self.assertFalse(form.is_valid())
        self.assertIn("marks", form.errors)

    def test_editing_marks_with_untouched_grade_regrades(self):
        result = ExamResult.objects.create(exam=self.exam, student=self.student,
                                           subject=self.subject, marks=45)  # A+
        self.assertEqual(result.grade, "A+")
        form = self._form("20", grade="A+", instance=result)  # marks changed, stale grade left in box
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["grade"], "C")  # 20/50 = 40%


class ExamResultBulkUploadGradingTests(TestCase):
    def setUp(self):
        self.admin = ExamResultAdmin(ExamResult, admin.site)
        self.klass = Class.objects.create(class_name="Class 7", academic_year=2026)
        self.subject = Subject.objects.create(class_obj=self.klass, subject_name="Physics")
        self.exam = Exam.objects.create(class_obj=self.klass, exam_name="Midterm",
                                        exam_date="2026-06-01", full_marks=100)
        self.student = Student.objects.create(student_code="PiC7B1001", class_obj=self.klass, full_name="T")
        self.headers = ["Student ID", "Full Name", "Physics"]

    def _upload(self, marks, exam=None):
        return self.admin._process_bulk_upload(
            _make_xlsx(self.headers, [["PiC7B1001", "T", marks]]), exam or self.exam)

    def test_uploaded_marks_get_a_grade(self):
        self._upload(85)
        self.assertEqual(ExamResult.objects.get().grade, "A+")

    def test_reupload_regrades(self):
        self._upload(85)
        self._upload(45)
        self.assertEqual(ExamResult.objects.get().grade, "C")

    def test_marks_above_full_marks_are_rejected_not_saved(self):
        result = self._upload(101)
        self.assertEqual(result["saved_count"], 0)
        self.assertTrue(any("full marks" in e for e in result["errors"]))
        self.assertFalse(ExamResult.objects.exists())

    def test_huge_marks_no_longer_crash_the_upload(self):
        result = self._upload(1000000)
        self.assertEqual(result["saved_count"], 0)
        self.assertTrue(result["errors"])

    def test_nan_and_infinity_are_rejected(self):
        for bad in ("NaN", "Infinity"):
            result = self._upload(bad)
            self.assertEqual(result["saved_count"], 0, bad)
            self.assertTrue(result["errors"], bad)

    def test_other_rows_still_saved_when_one_is_bad(self):
        other = Student.objects.create(student_code="PiC7B1002", class_obj=self.klass, full_name="U")
        upload = _make_xlsx(self.headers, [["PiC7B1001", "T", 500], ["PiC7B1002", "U", 70]])
        result = self.admin._process_bulk_upload(upload, self.exam)
        self.assertEqual(result["saved_count"], 1)
        self.assertEqual(ExamResult.objects.get().student, other)

    def test_marks_rounded_to_two_decimals(self):
        self._upload(84.555)
        self.assertEqual(str(ExamResult.objects.get().marks), "84.56")

    def test_exam_with_different_full_marks(self):
        exam = Exam.objects.create(class_obj=self.klass, exam_name="Quiz",
                                   exam_date="2026-06-10", full_marks=20)
        self.assertEqual(self._upload(17, exam)["saved_count"], 1)    # 85%
        self.assertEqual(ExamResult.objects.get(exam=exam).grade, "A+")
        self.assertEqual(self._upload(21, exam)["saved_count"], 0)    # > 20

    def test_duplicate_student_rows_are_flagged(self):
        upload = _make_xlsx(self.headers, [["PiC7B1001", "T", 40], ["PiC7B1001", "T", 90]])
        result = self.admin._process_bulk_upload(upload, self.exam)
        self.assertTrue(any("more than once" in e for e in result["errors"]))
        self.assertEqual(ExamResult.objects.get().marks, 90)


class StudentResultsDisplayTests(TestCase):
    def setUp(self):
        from django.contrib.auth import get_user_model
        self.user = get_user_model().objects.create_superuser("boss", "b@x.com", "pw12345!")
        self.client.force_login(self.user)
        self.klass = Class.objects.create(class_name="Class 9", academic_year=2026)
        self.exam = Exam.objects.create(class_obj=self.klass, exam_name="Half Yearly",
                                        exam_date="2026-06-01")
        self.student = Student.objects.create(
            student_code="PiC9B0001", class_obj=self.klass, full_name="Rina Akter",
            student_phone="01700000000")
        self.other = Student.objects.create(student_code="PiC9B0002", class_obj=self.klass,
                                            full_name="Someone Else")
        math = Subject.objects.create(class_obj=self.klass, subject_name="Mathematics")
        phys = Subject.objects.create(class_obj=self.klass, subject_name="Physics")
        ExamResult.objects.create(exam=self.exam, student=self.student, subject=math, marks=88)
        ExamResult.objects.create(exam=self.exam, student=self.student, subject=phys, marks=30)
        ExamResult.objects.create(exam=self.exam, student=self.other, subject=math, marks=99)

    def test_admin_student_page_shows_uploaded_results(self):
        response = self.client.get(reverse("admin:core_student_view", args=[self.student.pk]))
        self.assertEqual(response.status_code, 200)
        for text in ("Exam results", "Half Yearly", "Mathematics", "Physics", "Failed", "GPA"):
            self.assertContains(response, text)
        self.assertContains(response, "118 / 200")  # 88 + 30

    def test_admin_student_page_only_shows_that_students_results(self):
        response = self.client.get(reverse("admin:core_student_view", args=[self.student.pk]))
        reports = response.context["exam_reports"]
        self.assertEqual(len(reports), 1)
        self.assertEqual([r["marks"] for r in reports[0]["rows"]], [88, 30])  # not the other student's 99

    def test_admin_student_page_empty_state(self):
        student = Student.objects.create(student_code="PiC9B0003", class_obj=self.klass, full_name="New")
        response = self.client.get(reverse("admin:core_student_view", args=[student.pk]))
        self.assertContains(response, "No exam results recorded for this student yet.")

    def test_results_uploaded_in_bulk_appear_on_the_student_page(self):
        admin_obj = ExamResultAdmin(ExamResult, admin.site)
        exam2 = Exam.objects.create(class_obj=self.klass, exam_name="Final", exam_date="2026-12-01")
        Subject.objects.create(class_obj=self.klass, subject_name="Biology")
        admin_obj._process_bulk_upload(
            _make_xlsx(["Student ID", "Full Name", "Biology"], [["PiC9B0001", "Rina", 91]]), exam2)
        response = self.client.get(reverse("admin:core_student_view", args=[self.student.pk]))
        self.assertContains(response, "Final")
        self.assertContains(response, "Biology")
        self.assertEqual([r["exam"].exam_name for r in response.context["exam_reports"]],
                         ["Final", "Half Yearly"])

    def test_public_lookup_shows_totals_gpa_and_result(self):
        self.client.logout()
        response = self.client.post(reverse("results"),
                                    {"student_code": "PiC9B0001", "student_phone": "01700000000"})
        self.assertEqual(response.status_code, 200)
        for text in ("Half Yearly", "Mathematics", "118 / 200", "Failed", "Percentage"):
            self.assertContains(response, text)
        self.assertNotContains(response, "Someone Else")

    def test_recalculate_grades_fills_blank_and_fixes_stale_grades(self):
        ExamResult.objects.filter(student=self.student).update(grade=None)
        url = reverse("admin:core_exam_recalculate_grades", args=[self.exam.pk])
        self.client.get(url)  # GET must not change anything
        self.assertEqual(ExamResult.objects.filter(grade__isnull=True).count(), 2)
        response = self.client.post(url, follow=True)
        self.assertContains(response, "Grades recalculated")
        grades = dict(ExamResult.objects.filter(student=self.student).values_list("subject__subject_name", "grade"))
        self.assertEqual(grades, {"Mathematics": "A+", "Physics": "F"})

    def test_exam_pages_render_with_full_marks(self):
        self.assertContains(self.client.get(reverse("admin:core_exam_view", args=[self.exam.pk])), "Full marks 100")
        self.assertContains(self.client.get(reverse("admin:core_exam_change", args=[self.exam.pk])), "Full marks")

    def test_mark_sheet_template_has_readme_sheet_and_still_uploads(self):
        response = self.client.get(reverse("admin:core_examresult_bulk_template") + f"?exam={self.exam.pk}")
        workbook = openpyxl.load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ["Results", "Read me"])
        self.assertEqual([c.value for c in workbook["Results"][1]][:2], ["Student ID", "Full Name"])
        result = ExamResultAdmin(ExamResult, admin.site)._process_bulk_upload(BytesIO(response.content), self.exam)
        self.assertEqual(result["errors"], [])


class NavbarOrderTests(TestCase):
    """Notices sits in the main navbar, immediately before About, on both
    the desktop and the mobile menu."""

    def _menus(self):
        html = self.client.get(reverse("home")).content.decode()
        desktop = html[html.index('<nav class="hidden xl:flex'):html.index('<a href="/admin/" class="hidden xl:inline-flex')]
        mobile = html[html.index('<nav id="mobileMenu"'):html.index("</header>")]
        return desktop, mobile

    def _order(self, menu):
        import re
        return [re.sub(r"<[^>]+>", "", m).strip() for m in re.findall(r"<a [^>]*>.*?</a>", menu, re.S)
                if "/admin/" not in m]

    def test_desktop_order(self):
        desktop, _ = self._menus()
        main_bar = desktop[:desktop.index("More")]  # everything before the dropdown
        self.assertEqual(self._order(main_bar),
                         ["Home", "Classes", "Faculty", "Admission", "Results", "Notices", "About"])

    def test_notices_is_no_longer_inside_the_more_dropdown(self):
        desktop, _ = self._menus()
        self.assertEqual(self._order(desktop[desktop.index("More"):]),
                         ["Achievements", "Gallery", "Contact"])

    def test_mobile_order(self):
        _, mobile = self._menus()
        self.assertEqual(self._order(mobile), [
            "Home", "Classes", "Faculty", "Admission", "Results",
            "Notices", "About", "Achievements", "Gallery", "Contact", "Apply Now",
        ])


class ResponsiveLayoutTests(TestCase):
    """Guards for the mobile/tablet layout. (The pixel-level checks were
    done in a real browser; these stop the key pieces being removed.)"""

    def setUp(self):
        from django.contrib.auth import get_user_model
        self.client.force_login(get_user_model().objects.create_superuser("boss", "b@x.com", "pw12345!"))

    def test_every_standalone_template_declares_a_viewport(self):
        # Without it a phone renders the page at desktop width and shrinks it.
        import pathlib
        root = pathlib.Path(__file__).parent / "templates"
        missing = [str(p.relative_to(root)) for p in root.rglob("*.html")
                   if "<!DOCTYPE html>" in p.read_text(encoding="utf-8")
                   and 'name="viewport"' not in p.read_text(encoding="utf-8")]
        self.assertEqual(missing, [])

    def test_dashboard_has_a_mobile_menu(self):
        html = self.client.get(reverse("admin:index")).content.decode()
        for needle in ('id="menuToggle"', 'aria-controls="sidebar"', 'id="sidebar"',
                       'id="sidebarOverlay"', 'id="sidebarClose"'):
            self.assertIn(needle, html)

    def test_dashboard_tables_get_the_responsive_wrapper_script(self):
        html = self.client.get(reverse("admin:core_student_changelist")).content.decode()
        self.assertIn("table-scroll", html)
        self.assertIn("is-stacked", html)

    def test_month_navigators_use_the_phone_layout_class(self):
        for name in ("admin:core_studentpaymentreceipt_monthly_status", "admin:core_teachersalary_monthly_status"):
            try:
                url = reverse(name)
            except Exception:
                continue
            self.assertContains(self.client.get(url), "month-nav")

    def test_public_header_hamburger_and_apply_link_in_mobile_menu(self):
        html = self.client.get(reverse("home")).content.decode()
        self.assertIn('id="menuBtn"', html)
        mobile = html[html.index('<nav id="mobileMenu"'):html.index("</header>")]
        self.assertIn(reverse("apply"), mobile)

    def test_monthly_status_shows_the_amount_actually_paid(self):
        from datetime import date as _d
        klass = Class.objects.create(class_name="Class 6", academic_year=2026)
        student = Student.objects.create(student_code="PiC6B0001", class_obj=klass, full_name="Paid Student", status="Active")
        month = _d.today().replace(day=1)
        StudentPaymentReceipt.objects.create(
            student=student, payment_date=_d.today(), receipt_number="PAY-X-1", total_amount=3000, total_discount=0,
            items=[{"payment_type": "Monthly Fee", "payment_month": month.isoformat(), "label": "Monthly Fee", "amount": 3000}])
        response = self.client.get(reverse("admin:core_studentpaymentreceipt_monthly_status"))
        row = next(r for r in response.context["rows"] if r["student"].pk == student.pk)
        self.assertEqual(row["amount_paid"], 3000)